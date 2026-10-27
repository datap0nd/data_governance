"""Write the film's replay transcripts: plain-language business questions answered by Wizard's real tools over its
SYNTHETIC fixtures (and a synthetic targets workbook the person attaches). Every figure in the answer text is the
tool result it cites."""
import json, sys
from pathlib import Path
HERE = Path(__file__).resolve().parent
out = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / 'transcripts'
out.mkdir(parents=True, exist_ok=True)
Q3 = {"field": "fiscal_quarter", "values": ["2026-Q3"]}
Q23 = {"field": "fiscal_quarter", "values": ["2026-Q2", "2026-Q3"]}

quarter = {
  "id": "film-quarter-vs-targets", "kind": "ask", "data_mode": "SYNTHETIC",
  "match": ["against+targets", "against+target", "quarter+targets"],
  "steps": [
    {"say": "I'll take Q3 sell-out from GSCM, Samsung's share from ASAP and marketing spend from NERP, then compare sell-out with the targets in your file."},
    {"tool": "wizard_lookup_definitions", "args": {"query": "sell-out market share"}},
    {"tool": "gscm_run_report", "as": "sellout", "args": {"report_id": "gscm-sell-through-quarterly", "filters": [Q23],
      "group_by": ["market", "fiscal_quarter"], "measures": ["sell_out_units"]}},
    {"tool": "asap_run_report", "as": "share", "args": {"report_id": "asap-share-quarterly",
      "filters": [Q23, {"field": "brand", "values": ["Samsung"]}], "measures": ["volume_share_pct"]}},
    {"tool": "nerp_run_report", "as": "spend", "args": {"report_id": "nerp-mkt-spend-quarterly", "filters": [Q3],
      "group_by": ["market"], "measures": ["spend_usd"]}},
    {"tool": "wizard_query_attachment", "as": "targets", "args": {"file": "F1"}},
    {"tool": "wizard_calculate", "args": {"expression": "(q3 - q2) / q2 * 100", "variables": [
      {"name": "q3", "value": 1020000}, {"name": "q2", "value": 925000}]}},
    {"tool": "wizard_calculate", "args": {"expression": "actual / target * 100", "variables": [
      {"name": "actual", "value": 1020000}, {"name": "target", "value": 1000000}]}},
    {"tool": "wizard_render_visual", "as": "chart", "args": {"kind": "bar", "title": "Q3 sell-out against target",
      "subtitle": "Units, 2026-Q3. Targets from Q3 targets.xlsx.",
      "columns": [{"key": "market", "label": "Market"}, {"key": "target", "label": "Target", "type": "integer", "unit": "units"},
                  {"key": "sell_out", "label": "Sell-out", "type": "integer", "unit": "units"}],
      "rows": [["Saudi Arabia", 440000, 452000], ["UAE", 210000, 205000], ["Egypt", 250000, 262000], ["Morocco", 100000, 101000]],
      "x": "market", "series": ["target", "sell_out"], "evidence_ids": ["{sellout}", "{targets}"],
      "note": "Sell-out from GSCM; targets from the attached file."}},
    {"answer": "**A good quarter: sell-out rose 10.3% to 1.02 million units, 2% ahead of target.** Three of our four markets beat their target; the UAE is the one to watch.\n\n[{chart}]\n\n| Market | Q3 sell-out | vs Q2 | vs target | Samsung share | Q3 marketing spend |\n|---|---|---|---|---|---|\n| Saudi Arabia | 452,000 [{sellout}] | +10.2% | 103% [{targets}] | 32.0%, +0.8 pts [{share}] | $4.2M [{spend}] |\n| UAE | 205,000 [{sellout}] | +13.9% | 98% [{targets}] | 27.2%, −0.3 pts [{share}] | $2.6M [{spend}] |\n| Egypt | 262,000 [{sellout}] | +9.2% | 105% [{targets}] | 25.7%, +1.6 pts [{share}] | $1.5M [{spend}] |\n| Morocco | 101,000 [{sellout}] | +6.3% | 101% [{targets}] | 29.4%, +0.4 pts [{share}] | not posted [{spend}] |\n\n**What stands out**\n- **Saudi Arabia** added the most units, 42,000 more than in Q2, and finished 3% over target.\n- **Egypt** beat its target by the widest margin and gained the most share.\n- **UAE** is the one market below target, at 98%. Sell-out grew 13.9%, but Samsung's share slipped 0.3 points: the market grew faster than we did.\n- **Morocco** has no Q3 marketing spend posted in NERP yet.\n\nTargets come from your attached file, Q3 targets.xlsx, not from a company system."}
  ]}

stock = {
  "id": "film-channel-stock", "kind": "ask", "data_mode": "SYNTHETIC",
  "match": ["stock+building", "stock+piling", "channel+stock"],
  "steps": [
    {"say": "Stock builds up when we ship more into the channel than it sells, so I'll compare sell-in with sell-out for every model in Q3."},
    {"tool": "wizard_search_catalog", "args": {"query": "channel stock sell-in sell-out by model", "limit": 5}},
    {"tool": "gscm_run_report", "as": "quarterly", "args": {"report_id": "gscm-sell-through-quarterly", "filters": [Q23],
      "measures": ["closing_channel_stock_units", "sell_through_rate_pct"]}},
    {"tool": "gscm_run_report", "as": "models", "args": {"report_id": "gscm-sell-in-sell-out-monthly", "filters": [Q3],
      "group_by": ["market", "model"], "measures": ["sell_in_units", "sell_out_units"]}},
    {"tool": "gscm_run_report", "as": "stock", "args": {"report_id": "gscm-sell-in-sell-out-monthly",
      "filters": [{"field": "month", "values": ["2026-06", "2026-09"]}],
      "group_by": ["market", "model", "month"], "measures": ["channel_stock_units"]}},
    {"tool": "wizard_calculate", "args": {"expression": "(sep - jun) / jun * 100", "variables": [
      {"name": "sep", "value": 58675}, {"name": "jun", "value": 24092}]}},
    {"tool": "wizard_render_visual", "as": "chart", "args": {"kind": "bar", "title": "Where channel stock grew in Q3",
      "subtitle": "Channel stock at the end of June and September, units.",
      "columns": [{"key": "model", "label": "Model and market"}, {"key": "june", "label": "End of June", "type": "integer", "unit": "units"},
                  {"key": "september", "label": "End of September", "type": "integer", "unit": "units"}],
      "rows": [["Galaxy A37 · Egypt", 24092, 58675], ["Galaxy A57 · Saudi Arabia", 36620, 53254], ["Galaxy A17 · Morocco", 8267, 16044]],
      "x": "model", "series": ["june", "september"], "evidence_ids": ["{stock}"]}},
    {"answer": "**Egypt's Galaxy A37 is the one to act on: its channel stock more than doubled in Q3, from 24,092 to 58,675 units, because we shipped 34,583 more than the channel sold.**\n\n[{chart}]\n\n**Also rising**\n- **Galaxy A57 in Saudi Arabia:** up 45%, to 53,254 units [{stock}]. Sell-in ran 16,634 ahead of sell-out [{models}].\n- **Galaxy A17 in Morocco:** up 94%, to 16,044 units [{stock}].\n\n**Expected, not a concern**\n- The UAE's new Galaxy Z Flip8 and Z Fold8 hold their first launch stock (7,995 and 2,767 units).\n- Every other model moved by less than 3% [{stock}].\n\nEgypt's sell-through rate fell from 72.9% to 67.6% this quarter [{quarterly}]. **Suggested next step:** slow Galaxy A37 shipments to Egypt until the channel works through its stock."}
  ]}


def where(**kw):
    return [{"field": k, "values": [v]} for k, v in kw.items()]
SO, SH, SP, TG = "{E:gscm-sell-through-quarterly}", "{E:asap-share-quarterly}", "{E:nerp-mkt-spend-quarterly}", None
check = {
  "id": "film-quarter-check", "kind": "check", "data_mode": "SYNTHETIC",
  "match": ["against+targets", "against+target", "quarter+targets"],
  "steps": [
    {"say": "Re-checking the figures in the answer against the GSCM, ASAP and NERP reports they came from."},
    {"tool": "wizard_check_my_data", "as": "check", "args": {"claims": [
      {"label": "Total Q3 sell-out", "stated_value": 1020000, "evidence_id": SO, "measure": "sell_out_units",
       "where": where(fiscal_quarter="2026-Q3"), "period": "2026-Q3"},
      {"label": "Sell-out growth vs Q2 (%)", "stated_value": 10.3, "expression": "(q3 - q2) / q2 * 100", "inputs": [
        {"name": "q3", "evidence_id": SO, "measure": "sell_out_units", "where": where(fiscal_quarter="2026-Q3"), "period": "2026-Q3"},
        {"name": "q2", "evidence_id": SO, "measure": "sell_out_units", "where": where(fiscal_quarter="2026-Q2"), "period": "2026-Q2"}]},
      {"label": "Saudi Arabia Q3 sell-out", "stated_value": 452000, "evidence_id": SO, "measure": "sell_out_units",
       "where": where(market="SA", fiscal_quarter="2026-Q3"), "period": "2026-Q3"},
      {"label": "UAE sell-out growth (%)", "stated_value": 13.9, "expression": "(q3 - q2) / q2 * 100", "inputs": [
        {"name": "q3", "evidence_id": SO, "measure": "sell_out_units", "where": where(market="AE", fiscal_quarter="2026-Q3"), "period": "2026-Q3"},
        {"name": "q2", "evidence_id": SO, "measure": "sell_out_units", "where": where(market="AE", fiscal_quarter="2026-Q2"), "period": "2026-Q2"}]},
      {"label": "Egypt share gain (pts)", "stated_value": 1.6, "expression": "q3 - q2", "inputs": [
        {"name": "q3", "evidence_id": SH, "measure": "volume_share_pct", "where": where(market="EG", fiscal_quarter="2026-Q3"), "aggregate": "value", "period": "2026-Q3"},
        {"name": "q2", "evidence_id": SH, "measure": "volume_share_pct", "where": where(market="EG", fiscal_quarter="2026-Q2"), "aggregate": "value", "period": "2026-Q2"}]},
      {"label": "Saudi Arabia Q3 marketing spend", "stated_value": 4200000, "evidence_id": SP, "measure": "spend_usd",
       "where": where(market="SA"), "period": "2026-Q3"}]}},
    {"answer": "**All six figures I re-checked match the GSCM, ASAP and NERP reports they came from.** The answer stands: sell-out up 10.3% to 1.02 million units, led by Saudi Arabia, with the UAE the one market below target.\n\n- Total Q3 sell-out of 1,020,000 units and growth of 10.3% recompute from the GSCM rows.\n- Saudi Arabia's 452,000 units and $4.2M of spend, the UAE's 13.9% growth and Egypt's 1.6-point share gain match their source rows.\n\nThe targets come from your attached file, which no company system can confirm."}
  ]}

ST, TH = "{E:gscm-sell-in-sell-out-monthly}", "{E:gscm-sell-through-quarterly}"  # the latest is the month-end stock read
stock_rows = lambda market, model, month: {"evidence_id": ST, "measure": "channel_stock_units",
                                            "where": where(market=market, model=model, month=month), "period": month}
stock_check = {
  "id": "film-stock-check", "kind": "check", "data_mode": "SYNTHETIC",
  "match": ["stock+building", "stock+piling", "channel+stock"],
  "steps": [
    {"say": "Re-checking the figures in the answer against the GSCM reports they came from."},
    {"tool": "wizard_check_my_data", "as": "check", "args": {"claims": [
      {"label": "Egypt Galaxy A37 channel stock, end of June", "stated_value": 24092, **stock_rows("EG", "Galaxy A37", "2026-06")},
      {"label": "Egypt Galaxy A37 channel stock, end of September", "stated_value": 58675, **stock_rows("EG", "Galaxy A37", "2026-09")},
      {"label": "Egypt Galaxy A37 stock added in Q3 (units)", "stated_value": 34583, "expression": "sep - jun", "inputs": [
        {"name": "sep", **stock_rows("EG", "Galaxy A37", "2026-09")}, {"name": "jun", **stock_rows("EG", "Galaxy A37", "2026-06")}]},
      {"label": "Saudi Arabia Galaxy A57 stock growth (%)", "stated_value": 45, "expression": "(sep - jun) / jun * 100", "inputs": [
        {"name": "sep", **stock_rows("SA", "Galaxy A57", "2026-09")}, {"name": "jun", **stock_rows("SA", "Galaxy A57", "2026-06")}]},
      {"label": "Morocco Galaxy A17 stock growth (%)", "stated_value": 94, "expression": "(sep - jun) / jun * 100", "inputs": [
        {"name": "sep", **stock_rows("MA", "Galaxy A17", "2026-09")}, {"name": "jun", **stock_rows("MA", "Galaxy A17", "2026-06")}]},
      {"label": "Egypt Q3 sell-through rate (%)", "stated_value": 67.6, "evidence_id": TH, "measure": "sell_through_rate_pct",
       "where": where(market="EG", fiscal_quarter="2026-Q3"), "aggregate": "value", "period": "2026-Q3"}]}},
    {"answer": "**All six figures I re-checked match the GSCM reports they came from.** The answer stands: Egypt's Galaxy A37 channel stock more than doubled in Q3, from 24,092 to 58,675 units.\n\n- Egypt's Galaxy A37 stock at the end of June and September, and the 34,583 units it added, recompute from the GSCM rows.\n- The Galaxy A57's 45% rise in Saudi Arabia, the Galaxy A17's 94% rise in Morocco and Egypt's 67.6% sell-through rate match their source rows."}
  ]}

for t in (quarter, stock, check, stock_check):
    (out / f"{t['id']}.json").write_text(json.dumps(t, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
print("written", out)

# The attached file: synthetic Q3 sell-out targets by market (openpyxl, which Wizard's environment has).
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    sys.exit(0)
wb = Workbook(); ws = wb.active; ws.title = 'Targets'
ws.append(['market', 'market_name', 'fiscal_quarter', 'sell_out_target_units'])
for row in [('SA', 'Saudi Arabia', '2026-Q3', 440000), ('AE', 'UAE', '2026-Q3', 210000), ('EG', 'Egypt', '2026-Q3', 250000),
            ('MA', 'Morocco', '2026-Q3', 100000)]:
    ws.append(list(row))
for cell in ws[1]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='107C41')
for column, width in zip('ABCD', (10, 16, 16, 24)):
    ws.column_dimensions[column].width = width
for (cell,) in ws.iter_rows(min_row=2, min_col=4, max_col=4):
    cell.number_format = '#,##0'
wb.properties.creator = 'Wizard film (synthetic)'
(HERE / 'files').mkdir(exist_ok=True)
wb.save(HERE / 'files' / 'Q3 targets.xlsx')

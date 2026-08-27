#!/usr/bin/env python3
"""Build a fully offline Metronome sandbox in one deletable folder.

Creates everything the app needs to run locally with realistic fake data:

  <dest>/
    reports/            TMDL exports for six fake Power BI reports
    files/              random Excel/CSV files the reports point at
    usage/              Reports.csv / Users.csv / Report_views.csv (30 days)
    downloads/          fake flow run folders with downloaded artifacts
    pgdata/             a private PostgreSQL cluster (optional, see below)
    governance.db       Metronome SQLite DB pre-seeded with website flows,
                        people, tasks, docs, upstream systems, ...
    sandbox_config.json settings used by tools/run_sandbox.py
    README.txt          what this folder is and how to delete it

If PostgreSQL server binaries (initdb/pg_ctl) are found, a throwaway
cluster is created inside the folder with `track_commit_timestamp=on`,
a `metronome` database, and random tables plus a materialized view under
the `bi_reporting` schema, so SQL sources, probing, dependency discovery,
and Import Data all work end to end. Without them, use --skip-postgres:
Excel/CSV sources and everything else still work.

Usage:
    python tools/seed_sandbox.py                 # default: <repo>/local_sandbox
    python tools/seed_sandbox.py --force         # wipe and rebuild
    python tools/seed_sandbox.py --skip-postgres
    python tools/run_sandbox.py                  # then start the app

Deleting the sandbox: stop the app, run
    python tools/run_sandbox.py --stop
then delete the folder. Nothing outside the folder is modified.
"""

import argparse
import csv
import glob
import json
import os
import random
import shutil
import subprocess
import sys
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE_DIR))

PG_SUPERUSER = "postgres"
PG_APP_USER = "metronome"
PG_APP_PASSWORD = "metronome"
PG_DATABASE = "metronome"
PG_SCHEMA = "bi_reporting"

PEOPLE = [
    ("Rafael Cunha", "BI Lead", "rafael@example.com"),
    ("Ana Ferreira", "Data Analyst", "ana@example.com"),
    ("Miguel Santos", "Supply Chain Analyst", "miguel@example.com"),
    ("Sofia Costa", "Finance Controller", "sofia@example.com"),
    ("Pedro Alves", "Sales Ops", "pedro@example.com"),
]

REGIONS = ["EMEA", "APAC", "AMER"]
CARRIERS = ["DHL", "DB Schenker", "Kuehne+Nagel", "DSV", "Maersk"]
SEGMENTS = ["Enterprise", "SMB", "Distributor", "Retail"]
COUNTRIES = ["Portugal", "Spain", "Germany", "France", "Poland", "Brazil", "Mexico", "Japan"]
SKUS = [f"SKU-{n:05d}" for n in range(1001, 1041)]
KPIS = ["Revenue", "Gross Margin", "OTIF", "Forecast Accuracy", "Inventory Turns"]


def log(msg: str) -> None:
    print(f"[seed_sandbox] {msg}")


# ---------------------------------------------------------------------------
# Random tabular data
# ---------------------------------------------------------------------------

def _rand_dates(rng: random.Random, days_back: int, n: int) -> list[date]:
    today = date.today()
    return [today - timedelta(days=rng.randint(0, days_back)) for _ in range(n)]


def make_sales_rows(rng: random.Random, n: int = 1500) -> list[tuple]:
    rows = []
    for i, d in enumerate(_rand_dates(rng, 120, n), start=1):
        qty = rng.randint(1, 250)
        rows.append((
            i, d.isoformat(), rng.randint(10000, 10400), rng.choice(SKUS),
            qty, round(qty * rng.uniform(4.5, 90.0), 2), rng.choice(REGIONS),
        ))
    return rows


def make_kpi_rows(rng: random.Random) -> list[tuple]:
    rows = []
    today = date.today().replace(day=1)
    for m in range(12):
        month = (today - timedelta(days=30 * m)).replace(day=1)
        for kpi in KPIS:
            target = round(rng.uniform(60, 100), 1)
            rows.append((month.isoformat(), kpi, round(target * rng.uniform(0.8, 1.15), 1), target))
    return rows


def make_customer_rows(rng: random.Random, n: int = 400) -> list[tuple]:
    rows = []
    for i in range(n):
        signup = date.today() - timedelta(days=rng.randint(30, 2000))
        rows.append((
            10000 + i, f"Customer {10000 + i}", rng.choice(SEGMENTS),
            rng.choice(COUNTRIES), signup.isoformat(), rng.random() > 0.15,
        ))
    return rows


def make_shipment_rows(rng: random.Random, n: int = 900) -> list[tuple]:
    rows = []
    for i, d in enumerate(_rand_dates(rng, 90, n), start=1):
        rows.append((
            i, d.isoformat(), rng.choice(CARRIERS), rng.choice(COUNTRIES),
            rng.choice(COUNTRIES), rng.random() > 0.12, rng.random() > 0.08,
        ))
    return rows


def make_gl_rows(rng: random.Random, n: int = 1200) -> list[tuple]:
    accounts = ["4000-Revenue", "5000-COGS", "6100-Freight", "6200-Warehousing", "7000-Overhead"]
    rows = []
    for i, d in enumerate(_rand_dates(rng, 365, n), start=1):
        rows.append((
            i, d.isoformat(), rng.choice(accounts), f"CC-{rng.randint(100, 130)}",
            round(rng.uniform(-50000, 90000), 2),
        ))
    return rows


# ---------------------------------------------------------------------------
# Excel / CSV files
# ---------------------------------------------------------------------------

def write_excel(path: Path, sheet: str, headers: list[str], rows: list[tuple]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    wb.save(path)


def write_csv(path: Path, headers: list[str], rows: list[tuple]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        writer.writerows(rows)


def build_files(dest: Path, rng: random.Random) -> None:
    files = dest / "files"
    files.mkdir(parents=True, exist_ok=True)

    write_excel(
        files / "sku_master.xlsx", "SKUs",
        ["SKU", "Description", "Category", "UnitCost"],
        [(sku, f"Component {sku[-5:]}", rng.choice(["Display", "Battery", "PCB", "Housing"]),
          round(rng.uniform(0.4, 40), 2)) for sku in SKUS],
    )
    write_excel(
        files / "inventory_snapshot.xlsx", "Inventory",
        ["SKU", "Warehouse", "OnHand", "Allocated", "SnapshotDate"],
        [(rng.choice(SKUS), f"WH-{rng.randint(1, 6)}", rng.randint(0, 5000),
          rng.randint(0, 800), date.today().isoformat()) for _ in range(600)],
    )
    write_excel(
        files / "carrier_rates.xlsx", "Rates",
        ["Carrier", "Lane", "RatePerKg", "ValidFrom"],
        [(c, f"{rng.choice(COUNTRIES)} -> {rng.choice(COUNTRIES)}",
          round(rng.uniform(0.8, 6.5), 2), "2026-01-01") for c in CARRIERS for _ in range(6)],
    )
    write_excel(
        files / "gl_extract.xlsx", "GL",
        ["EntryId", "PostingDate", "Account", "CostCenter", "AmountEUR"],
        make_gl_rows(rng, 400),
    )
    write_csv(
        files / "warehouse_locations.csv",
        ["Warehouse", "City", "Country", "Capacity"],
        [(f"WH-{i}", rng.choice(["Porto", "Lisbon", "Madrid", "Lyon", "Gdansk", "Monterrey"]),
          rng.choice(COUNTRIES), rng.randint(5000, 60000)) for i in range(1, 7)],
    )
    write_csv(
        files / "churn_scores.csv",
        ["CustomerId", "ChurnScore", "ScoredAt"],
        [(10000 + i, round(rng.random(), 3), date.today().isoformat()) for i in range(400)],
    )
    log("files/: 4 Excel workbooks + 2 CSVs written")


# ---------------------------------------------------------------------------
# TMDL report exports
# ---------------------------------------------------------------------------

def _tmdl_partition(table: str, m_lines: list[str]) -> str:
    body = "\n".join(f"\t\t\t{line}" for line in m_lines)
    name = f"'{table}'" if " " in table else table
    return f"\tpartition {name} = m\n\t\tmode: import\n\t\tsource =\n{body}\n"


def _tmdl_table(table: str, columns: list[str], m_lines: list[str],
                measures: list[tuple[str, str]] | None = None) -> str:
    name = f"'{table}'" if " " in table else table
    out = [f"table {name}", ""]
    for col in columns:
        cname = f"'{col}'" if " " in col else col
        out.append(f"\tcolumn {cname}")
    out.append("")
    for mname, dax in measures or []:
        out.append(f"\tmeasure '{mname}' = {dax}")
    if measures:
        out.append("")
    out.append(_tmdl_partition(table, m_lines))
    return "\n".join(out) + "\n"


def _m_postgres(pg_host_port: str, schema: str, table: str) -> list[str]:
    return [
        "let",
        f'\tSource = PostgreSQL.Database("{pg_host_port}", "{PG_DATABASE}"),',
        f'\tData = Source{{[Schema="{schema}",Item="{table}"]}}[Data]',
        "in",
        "\tData",
    ]


def _m_excel(path: Path, sheet: str) -> list[str]:
    return [
        "let",
        f'\tSource = Excel.Workbook(File.Contents("{path}"), null, true),',
        f'\tData = Source{{[Item="{sheet}",Kind="Sheet"]}}[Data]',
        "in",
        "\tData",
    ]


def _m_csv(path: Path) -> list[str]:
    return [
        "let",
        f'\tSource = Csv.Document(File.Contents("{path}"),[Delimiter=",", Encoding=65001]),',
        "\tData = Table.PromoteHeaders(Source)",
        "in",
        "\tData",
    ]


def _m_owner(name: str) -> list[str]:
    return [f'#table({{"Value"}}, {{{{"{name}"}}}})']


def build_reports(dest: Path, pg_host_port: str) -> list[str]:
    files = dest / "files"
    reports = {
        "Weekly_Sales": {
            "business_owner": "Pedro Alves", "report_owner": "Rafael Cunha",
            "tables": [
                ("Main", ["OrderId", "OrderDate", "CustomerId", "SKU", "Qty", "Amount", "Region"],
                 _m_postgres(pg_host_port, PG_SCHEMA, "sales_orders"),
                 [("Total Sales", "SUM(Main[Amount])"), ("Total Units", "SUM(Main[Qty])")]),
                ("SKU Master", ["SKU", "Description", "Category", "UnitCost"],
                 _m_excel(files / "sku_master.xlsx", "SKUs"), None),
            ],
        },
        "Monthly_KPI": {
            "business_owner": "Sofia Costa", "report_owner": "Rafael Cunha",
            "tables": [
                ("KPIs", ["Month", "KpiName", "Value", "Target"],
                 _m_postgres(pg_host_port, PG_SCHEMA, "kpi_monthly"),
                 [("Attainment", "DIVIDE(SUM(KPIs[Value]), SUM(KPIs[Target]))")]),
                ("KPI Summary", ["KpiName", "AvgValue"],
                 _m_postgres(pg_host_port, PG_SCHEMA, "mv_kpi_summary"), None),
            ],
        },
        "Inventory_Health": {
            "business_owner": "Miguel Santos", "report_owner": "Ana Ferreira",
            "tables": [
                ("Inventory", ["SKU", "Warehouse", "OnHand", "Allocated", "SnapshotDate"],
                 _m_excel(files / "inventory_snapshot.xlsx", "Inventory"), None),
                ("Warehouses", ["Warehouse", "City", "Country", "Capacity"],
                 _m_csv(files / "warehouse_locations.csv"), None),
            ],
        },
        "Customer_Churn": {
            "business_owner": "Pedro Alves", "report_owner": "Ana Ferreira",
            "tables": [
                ("Customers", ["CustomerId", "Name", "Segment", "Country", "SignupDate", "Active"],
                 _m_postgres(pg_host_port, PG_SCHEMA, "customers"), None),
                ("Churn Scores", ["CustomerId", "ChurnScore", "ScoredAt"],
                 _m_csv(files / "churn_scores.csv"), None),
            ],
        },
        "Logistics_OTIF": {
            "business_owner": "Miguel Santos", "report_owner": "Rafael Cunha",
            "tables": [
                ("Shipments", ["ShipmentId", "ShipDate", "Carrier", "Origin", "Destination", "OnTime", "InFull"],
                 _m_postgres(pg_host_port, PG_SCHEMA, "shipments"),
                 [("OTIF %", "DIVIDE(COUNTROWS(FILTER(Shipments, Shipments[OnTime] && Shipments[InFull])), COUNTROWS(Shipments))")]),
                ("Carrier Rates", ["Carrier", "Lane", "RatePerKg", "ValidFrom"],
                 _m_excel(files / "carrier_rates.xlsx", "Rates"), None),
            ],
        },
        "Finance_PL": {
            "business_owner": "Sofia Costa", "report_owner": "Sofia Costa",
            "tables": [
                ("GL Entries", ["EntryId", "PostingDate", "Account", "CostCenter", "AmountEUR"],
                 _m_postgres(pg_host_port, PG_SCHEMA, "gl_entries"), None),
                ("GL Extract", ["EntryId", "PostingDate", "Account", "CostCenter", "AmountEUR"],
                 _m_excel(files / "gl_extract.xlsx", "GL"), None),
            ],
        },
    }

    for name, spec in reports.items():
        tables_dir = dest / "reports" / name / f"{name}.SemanticModel" / "Definition" / "Tables"
        tables_dir.mkdir(parents=True, exist_ok=True)
        for table, columns, m_lines, measures in spec["tables"]:
            (tables_dir / f"{table}.tmdl").write_text(
                _tmdl_table(table, columns, m_lines, measures), encoding="utf-8"
            )
        (tables_dir / "Business Owner.tmdl").write_text(
            _tmdl_table("Business Owner", ["Value"], _m_owner(spec["business_owner"])),
            encoding="utf-8",
        )
        (tables_dir / "Report Owner.tmdl").write_text(
            _tmdl_table("Report Owner", ["Value"], _m_owner(spec["report_owner"])),
            encoding="utf-8",
        )
    log(f"reports/: {len(reports)} TMDL report exports written")
    return list(reports)


# ---------------------------------------------------------------------------
# Usage CSVs
# ---------------------------------------------------------------------------

def build_usage(dest: Path, rng: random.Random, report_names: list[str]) -> None:
    usage = dest / "usage"
    usage.mkdir(parents=True, exist_ok=True)

    guids = {name: str(uuid.uuid4()) for name in report_names}
    users = [(f"key-{i}", person[2]) for i, person in enumerate(PEOPLE)]
    users += [(f"key-{i}", f"viewer{i}@example.com") for i in range(len(PEOPLE), 14)]

    write_csv(usage / "Reports.csv", ["ReportGuid", "ReportName"],
              [(guid, name) for name, guid in guids.items()])
    write_csv(usage / "Users.csv", ["UserKey", "UserId"], users)

    views = []
    today = date.today()
    for day_back in range(30):
        d = today - timedelta(days=day_back)
        for name, guid in guids.items():
            for user_key, _ in users:
                for _ in range(rng.choices([0, 1, 2], weights=[62, 28, 10])[0]):
                    views.append((d.isoformat(), guid, name, user_key))
    rng.shuffle(views)
    write_csv(usage / "Report_views.csv", ["Date", "ReportId", "ReportName", "UserKey"], views)
    log(f"usage/: {len(views)} report views across 30 days written")


# ---------------------------------------------------------------------------
# PostgreSQL sandbox cluster
# ---------------------------------------------------------------------------

def find_pg_bin() -> Path | None:
    for tool in ("initdb",):
        found = shutil.which(tool)
        if found:
            return Path(found).parent
    candidates = sorted(glob.glob("/usr/lib/postgresql/*/bin")) + \
        sorted(glob.glob(r"C:\Program Files\PostgreSQL\*\bin"))
    for cand in reversed(candidates):
        if (Path(cand) / "initdb").exists() or (Path(cand) / "initdb.exe").exists():
            return Path(cand)
    return None


def _pg_runner() -> list[str]:
    """PostgreSQL refuses to run as root; wrap commands with runuser then."""
    if os.name != "nt" and os.geteuid() == 0:
        import pwd
        try:
            pwd.getpwnam(PG_SUPERUSER)
            return ["runuser", "-u", PG_SUPERUSER, "--"]
        except KeyError:
            raise SystemExit(
                "Running as root but no 'postgres' system user exists to own the "
                "cluster. Re-run as a normal user or with --skip-postgres."
            )
    return []


def pg_is_running(pg_bin: Path, pgdata: Path) -> bool:
    res = subprocess.run(
        _pg_runner() + [str(pg_bin / "pg_ctl"), "-D", str(pgdata), "status"],
        capture_output=True, text=True,
    )
    return res.returncode == 0


def start_postgres(pg_bin: Path, pgdata: Path) -> None:
    if pg_is_running(pg_bin, pgdata):
        return
    subprocess.run(
        _pg_runner() + [str(pg_bin / "pg_ctl"), "-D", str(pgdata),
                        "-l", str(pgdata / "postgres.log"), "-w", "start"],
        check=True,
    )


def stop_postgres(pg_bin: Path, pgdata: Path) -> None:
    if pgdata.exists() and pg_is_running(pg_bin, pgdata):
        subprocess.run(
            _pg_runner() + [str(pg_bin / "pg_ctl"), "-D", str(pgdata), "-m", "fast", "stop"],
            check=False,
        )


def build_postgres(dest: Path, rng: random.Random, pg_bin: Path, port: int) -> None:
    pgdata = dest / "pgdata"
    runner = _pg_runner()

    pgdata.mkdir(parents=True, exist_ok=True)
    if runner:
        shutil.chown(pgdata, user=PG_SUPERUSER, group=PG_SUPERUSER)
    pgdata.chmod(0o700)

    subprocess.run(
        runner + [str(pg_bin / "initdb"), "-D", str(pgdata), "-U", PG_SUPERUSER,
                  "-E", "UTF8", "-A", "trust"],
        check=True, capture_output=True, text=True,
    )
    with (pgdata / "postgresql.conf").open("a", encoding="utf-8") as fh:
        fh.write(
            "\n# --- Metronome sandbox overrides ---\n"
            f"port = {port}\n"
            "listen_addresses = '127.0.0.1'\n"
            f"unix_socket_directories = '{pgdata}'\n"
            "track_commit_timestamp = on\n"
        )
    start_postgres(pg_bin, pgdata)

    import psycopg2

    admin = psycopg2.connect(host="127.0.0.1", port=port, user=PG_SUPERUSER, dbname="postgres")
    admin.autocommit = True
    with admin.cursor() as cur:
        cur.execute(
            f"CREATE ROLE {PG_APP_USER} LOGIN PASSWORD %s CREATEDB", (PG_APP_PASSWORD,)
        )
        cur.execute(f"CREATE DATABASE {PG_DATABASE} OWNER {PG_APP_USER}")
    admin.close()

    conn = psycopg2.connect(host="127.0.0.1", port=port, user=PG_APP_USER,
                            password=PG_APP_PASSWORD, dbname=PG_DATABASE)
    conn.autocommit = True
    with conn.cursor() as cur:
        cur.execute(f"CREATE SCHEMA {PG_SCHEMA}")
        cur.execute(f"""CREATE TABLE {PG_SCHEMA}.sales_orders (
            order_id INT PRIMARY KEY, order_date DATE, customer_id INT,
            sku TEXT, qty INT, amount NUMERIC(12,2), region TEXT)""")
        cur.execute(f"""CREATE TABLE {PG_SCHEMA}.kpi_monthly (
            month DATE, kpi_name TEXT, value NUMERIC(8,1), target NUMERIC(8,1))""")
        cur.execute(f"""CREATE TABLE {PG_SCHEMA}.customers (
            customer_id INT PRIMARY KEY, name TEXT, segment TEXT,
            country TEXT, signup_date DATE, active BOOLEAN)""")
        cur.execute(f"""CREATE TABLE {PG_SCHEMA}.shipments (
            shipment_id INT PRIMARY KEY, ship_date DATE, carrier TEXT,
            origin TEXT, destination TEXT, on_time BOOLEAN, in_full BOOLEAN)""")
        cur.execute(f"""CREATE TABLE {PG_SCHEMA}.gl_entries (
            entry_id INT PRIMARY KEY, posting_date DATE, account TEXT,
            cost_center TEXT, amount_eur NUMERIC(14,2))""")

        cur.executemany(
            f"INSERT INTO {PG_SCHEMA}.sales_orders VALUES (%s,%s,%s,%s,%s,%s,%s)",
            make_sales_rows(rng))
        cur.executemany(
            f"INSERT INTO {PG_SCHEMA}.kpi_monthly VALUES (%s,%s,%s,%s)",
            make_kpi_rows(rng))
        cur.executemany(
            f"INSERT INTO {PG_SCHEMA}.customers VALUES (%s,%s,%s,%s,%s,%s)",
            make_customer_rows(rng))
        cur.executemany(
            f"INSERT INTO {PG_SCHEMA}.shipments VALUES (%s,%s,%s,%s,%s,%s,%s)",
            make_shipment_rows(rng))
        cur.executemany(
            f"INSERT INTO {PG_SCHEMA}.gl_entries VALUES (%s,%s,%s,%s,%s)",
            make_gl_rows(rng))

        cur.execute(f"""CREATE MATERIALIZED VIEW {PG_SCHEMA}.mv_kpi_summary AS
            SELECT kpi_name, ROUND(AVG(value), 1) AS avg_value
            FROM {PG_SCHEMA}.kpi_monthly GROUP BY kpi_name""")
    conn.close()
    log(f"pgdata/: PostgreSQL cluster on 127.0.0.1:{port}, db '{PG_DATABASE}', "
        f"5 tables + 1 materialized view in schema '{PG_SCHEMA}'")


# ---------------------------------------------------------------------------
# governance.db (flows, people, tasks, docs, ...)
# ---------------------------------------------------------------------------

def build_governance_db(dest: Path, rng: random.Random) -> None:
    os.environ["DG_DB_PATH"] = str(dest / "governance.db")
    import importlib

    import app.config as app_config
    importlib.reload(app_config)
    import app.database as app_database
    importlib.reload(app_database)
    app_database.init_db()

    now = datetime.now(timezone.utc)
    iso = now.isoformat()
    downloads = dest / "downloads"

    with app_database.get_db() as db:
        for name, role, email in PEOPLE:
            db.execute(
                "INSERT INTO people (name, role, email, include_all_alerts) VALUES (?,?,?,0)",
                (name, role, email),
            )
        db.execute(
            "INSERT INTO user_ips (ip_address, person_name, hostname) VALUES ('127.0.0.1', ?, 'sandbox-pc')",
            (PEOPLE[0][0],),
        )

        for name, code, day in [("SAP ECC", "SAP", "Monday"), ("ASAP Portal", "ASAP", "Saturday"),
                                ("GSCM", "GSCM", "Sunday")]:
            db.execute(
                "INSERT INTO upstream_systems (name, code, refresh_day, discovered_by) VALUES (?,?,?,'manual')",
                (name, code, day),
            )

        for title, status, priority, assignee in [
            ("Review stale sources after sandbox seed", "backlog", "medium", PEOPLE[0][0]),
            ("Tune churn model thresholds", "in_progress", "high", PEOPLE[1][0]),
            ("Document OTIF calculation", "done", "low", PEOPLE[2][0]),
        ]:
            db.execute(
                "INSERT INTO tasks (title, status, priority, assigned_to) VALUES (?,?,?,?)",
                (title, status, priority, assignee),
            )

        db.execute(
            """INSERT INTO power_automate_flows (name, description, owner, schedule, status, account)
               VALUES ('Refresh churn extract', 'Copies churn_scores.csv from the DS share',
                       ?, 'Daily 06:00', 'active', 'svc-flows@example.com')""",
            (PEOPLE[1][0],),
        )
        db.execute(
            """INSERT INTO custom_reports (name, description, frequency, owner, stakeholders, status)
               VALUES ('Quarterly Business Review pack', 'Manual PowerPoint refresh from Monthly_KPI',
                       'quarterly', ?, 'Leadership team', 'active')""",
            (PEOPLE[3][0],),
        )
        db.execute(
            """INSERT INTO documentation (title, business_purpose, business_audience, business_cadence,
                                          technical_sources, status, created_by)
               VALUES ('Weekly Sales pipeline', 'Tracks sell-out orders by region and SKU.',
                       'Sales Ops, Leadership', 'Weekly, Monday 08:00',
                       'bi_reporting.sales_orders (PostgreSQL), sku_master.xlsx', 'published', ?)""",
            (PEOPLE[0][0],),
        )
        for email in ("rafael@example.com", "sofia@example.com"):
            db.execute("INSERT INTO premium_viewers (email) VALUES (?)", (email,))

        # --- Website flow catalog (what a portal discovery scan would find) ---
        db.execute(
            """INSERT INTO flow_sites
               (name, adapter, base_url, auth_url, discovery_enabled, discovery_interval_hours,
                discovery_scope_json, discovery_weekday, discovery_time,
                last_scan_at, last_scan_status, enabled)
               VALUES ('ASAP Portal', 'web_export', 'https://asap.example.com/portal',
                       'https://sso.example.com/login', 1, 168, '["Mobile", "Supply Chain"]',
                       'saturday', '06:00', ?, 'succeeded', 1)""",
            (iso,),
        )
        site_id = db.execute("SELECT id FROM flow_sites WHERE name='ASAP Portal'").fetchone()["id"]

        catalog = [
            ("Mobile > Sales > Weekly Sell-Out",
             "https://asap.example.com/portal/mobile/sales/weekly-sellout",
             [{"label": "Raw Data Export"}, {"label": "Summary Export"}],
             [("week", "Week", "Week selector", "week", [], 1),
              ("region", "Region", "Region prompt", "select", REGIONS, 1),
              ("channel", "Channel", "Channel prompt", "multi_select",
               ["Online", "Retail", "Distributor"], 0)]),
            ("Mobile > Sales > Sell-In by Account",
             "https://asap.example.com/portal/mobile/sales/sell-in",
             [{"label": "Raw Data Export"}],
             [("week", "Week", "Week selector", "week", [], 1),
              ("account", "Account", "Account prompt", "select",
               ["All", "Key Accounts", "Open Market"], 0)]),
            ("Supply Chain > M Tracker",
             "https://asap.example.com/portal/sc/m-tracker",
             [{"label": "Dashboard Download"}],
             [("scope", "Scope", "Scope tabs", "select", ["Global", "Europe"], 0)]),
        ]
        for name, url, views, filters in catalog:
            db.execute(
                """INSERT INTO flow_reports
                   (site_id, name, report_url, download_text, automation_json,
                    discovery_key, source_kind, last_seen_at, stale, enabled)
                   VALUES (?,?,?,'Export', ?, ?, 'discovered', ?, 0, 1)""",
                (site_id, name, url, json.dumps({"export_views": views}),
                 name.lower().replace(" ", "-"), iso),
            )
            report_id = db.execute(
                "SELECT id FROM flow_reports WHERE site_id=? AND name=?", (site_id, name)
            ).fetchone()["id"]
            for pos, (key, label, control_label, ctype, options, required) in enumerate(filters):
                db.execute(
                    """INSERT INTO flow_report_filters
                       (report_id, filter_key, label, control_label, control_type,
                        options_json, required, source_kind, last_seen_at, stale, position, enabled)
                       VALUES (?,?,?,?,?,?,?, 'discovered', ?, 0, ?, 1)""",
                    (report_id, key, label, control_label, ctype,
                     json.dumps(options), required, iso, pos),
                )

        # --- Two saved flows plus believable run history ---
        this_week = date.today().isocalendar()
        week_key = f"{this_week[0]}{this_week[1]:02d}"
        flows = [
            ("Weekly Sell-Out EMEA", "Mobile > Sales > Weekly Sell-Out",
             ["Raw Data Export"], {"region": "EMEA", "channel": ["Online", "Retail"]},
             "weekly", "monday", "07:30", "sellout_emea_{week}.csv"),
            ("M Tracker snapshot", "Supply Chain > M Tracker",
             ["Dashboard Download"], {"scope": "Global"},
             "daily", None, "06:45", "m_tracker_{date}.xlsx"),
        ]
        for name, report_name, views, selections, sched_type, sched_days, sched_time, template in flows:
            report_id = db.execute(
                "SELECT id FROM flow_reports WHERE site_id=? AND name=?", (site_id, report_name)
            ).fetchone()["id"]
            target = downloads / name.lower().replace(" ", "_")
            db.execute(
                """INSERT INTO flows
                   (name, source_type, site_id, report_id, export_views_json, enabled,
                    selections_json, download_mode, period_strategy, file_format, browser_mode,
                    start_week, end_week, target_folder, filename_template,
                    schedule_type, schedule_time, schedule_days,
                    last_run_at, last_success_at, last_status, owner_person_id, created_by)
                   VALUES (?,?,?,?,?,0,?,?,'fixed',?, 'headless', ?, ?, ?, ?, ?, ?, ?, ?, ?, 'succeeded',
                           (SELECT id FROM people WHERE name=?), ?)""",
                (name, "portal", site_id, report_id, json.dumps(views),
                 json.dumps(selections), "single",
                 "csv" if template.endswith("csv") else "xlsx",
                 week_key, week_key, str(target), template,
                 sched_type, sched_time, json.dumps([sched_days] if sched_days else []),
                 iso, iso, PEOPLE[0][0], PEOPLE[0][0]),
            )
            flow_id = db.execute("SELECT id FROM flows WHERE name=?", (name,)).fetchone()["id"]

            for run_no in range(1, 4):
                started = now - timedelta(days=(3 - run_no) * 7 + 1)
                run_date = started.strftime("%d-%m-%Y")
                run_folder = target / f"#{run_no}_{run_date}"
                run_folder.mkdir(parents=True, exist_ok=True)
                filename = template.replace("{week}", week_key).replace(
                    "{date}", started.strftime("%Y%m%d"))
                artifact = run_folder / filename
                if filename.endswith(".xlsx"):
                    write_excel(artifact, "Data", ["Metric", "Value"],
                                [(f"Metric {i}", rng.randint(1, 999)) for i in range(20)])
                else:
                    write_csv(artifact, ["Week", "Region", "SKU", "Units"],
                              [(week_key, rng.choice(REGIONS), rng.choice(SKUS),
                                rng.randint(1, 500)) for _ in range(150)])
                status = "succeeded" if run_no != 2 else "failed"
                error = None if status == "succeeded" else \
                    "Portal returned HTTP 500 while opening the export view (retry also failed)"
                db.execute(
                    """INSERT INTO flow_runs
                       (flow_id, trigger_type, status, requested_by, worker_id, job_json,
                        error, created_at, claimed_at, started_at, finished_at)
                       VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
                    (flow_id, "schedule" if run_no > 1 else "manual", status, PEOPLE[0][0],
                     "sandbox-worker",
                     json.dumps({"flow": name, "export_views": views, "selections": selections}),
                     error, started.isoformat(), started.isoformat(), started.isoformat(),
                     (started + timedelta(minutes=6)).isoformat()),
                )
                run_id = db.execute("SELECT MAX(id) AS id FROM flow_runs").fetchone()["id"]
                stages = [("running", "navigation", "Signed in and opened the report"),
                          ("running", "configuration", "Applied filters"),
                          ("running", "export", "Requested export")]
                if status == "succeeded":
                    stages.append(("succeeded", "transfer", f"Downloaded {filename}"))
                else:
                    stages.append(("failed", "export", error))
                for st, stage, message in stages:
                    db.execute(
                        """INSERT INTO flow_run_events (run_id, status, stage, message, created_at)
                           VALUES (?,?,?,?,?)""",
                        (run_id, st, stage, message, started.isoformat()),
                    )
                if status == "succeeded":
                    db.execute(
                        """INSERT INTO flow_run_files
                           (run_id, period_key, file_path, filename, file_size, row_count, status)
                           VALUES (?,?,?,?,?,?, 'downloaded')""",
                        (run_id, week_key, str(artifact), filename,
                         artifact.stat().st_size, 150, ),
                    )
                for phase, ms in [("navigation", rng.randint(8000, 20000)),
                                  ("configuration", rng.randint(3000, 9000)),
                                  ("export", rng.randint(15000, 60000))]:
                    db.execute(
                        """INSERT INTO flow_operation_timings
                           (operation_type, phase, run_id, site_id, report_id, duration_ms, status)
                           VALUES ('download', ?, ?, ?, ?, ?, ?)""",
                        (phase, run_id, site_id, report_id, ms,
                         "succeeded" if status == "succeeded" else "failed"),
                    )

        db.execute(
            """INSERT INTO flow_workers (worker_id, display_name, capabilities_json, status, last_seen_at)
               VALUES ('sandbox-worker', 'Sandbox BI desktop', '{"headless": true}', 'offline', ?)""",
            (iso,),
        )
    log("governance.db: people, tasks, upstream systems, website flow catalog, "
        "2 flows with run history, docs seeded")


# ---------------------------------------------------------------------------
# First scan + probe, so the app opens with a populated dashboard
# ---------------------------------------------------------------------------

def bootstrap_app_state(dest: Path) -> None:
    """Run the real scanner pipeline once against the fresh sandbox.

    Discovers the TMDL reports, registers sources, applies freshness rules
    (one source is made deliberately stale so alerts/actions show up),
    probes freshness, and imports the usage CSVs.
    """
    import app.database as app_database
    from app.scanner import runner, prober
    from app.usage import sync_usage_from_csv

    result = runner.run_scan(reports_path=str(dest), run_followup_probe=False)
    log(f"scan: {result.get('reports_scanned')} reports, "
        f"{result.get('sources_found')} sources registered")

    # One deliberately stale file source, so the dashboard shows a real alert.
    stale_file = dest / "files" / "churn_scores.csv"
    old = (datetime.now(timezone.utc) - timedelta(days=10)).timestamp()
    os.utime(stale_file, (old, old))

    with app_database.get_db() as db:
        db.execute(
            "UPDATE sources SET freshness_rule_type='daily' WHERE type='postgresql'")
        db.execute(
            """UPDATE sources SET freshness_rule_type='custom', custom_fresh_days=14
               WHERE type='excel'""")
        db.execute(
            """UPDATE sources SET freshness_rule_type='daily', custom_fresh_days=NULL
               WHERE name='churn_scores.csv'""")

    probe = prober.run_probe()
    log(f"probe: {probe.get('statuses')}")

    with app_database.get_db() as db:
        usage = sync_usage_from_csv(db, force=True)
    log(f"usage: {usage.get('message')}")


# ---------------------------------------------------------------------------
# main
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dest", default=str(BASE_DIR / "local_sandbox"),
                        help="Sandbox folder (default: <repo>/local_sandbox)")
    parser.add_argument("--pg-port", type=int, default=5433,
                        help="Port for the sandbox PostgreSQL (default: 5433)")
    parser.add_argument("--skip-postgres", action="store_true",
                        help="Do not create the PostgreSQL cluster")
    parser.add_argument("--force", action="store_true",
                        help="Delete an existing sandbox folder first")
    parser.add_argument("--seed", type=int, default=None,
                        help="Random seed (default: random each run)")
    args = parser.parse_args()

    dest = Path(args.dest).resolve()
    rng = random.Random(args.seed)

    pg_bin = None if args.skip_postgres else find_pg_bin()
    if not args.skip_postgres and pg_bin is None:
        log("WARNING: PostgreSQL binaries not found - continuing with --skip-postgres behavior")

    if dest.exists() and any(dest.iterdir()):
        if not args.force:
            raise SystemExit(f"{dest} already exists. Use --force to wipe and rebuild it.")
        if pg_bin is not None:
            stop_postgres(pg_bin, dest / "pgdata")
        shutil.rmtree(dest)
    dest.mkdir(parents=True, exist_ok=True)

    # Point the app modules used below (and any child imports) at the sandbox
    # before anything from app/ is imported.
    os.environ.update({
        "DG_DB_PATH": str(dest / "governance.db"),
        "DG_TMDL_ROOT": str(dest),
        "usage_files_path": str(dest / "usage"),
    })
    if pg_bin is not None:
        os.environ.update({
            "PGHOST": "127.0.0.1",
            "PGPORT": str(args.pg_port),
            "PGUSER": PG_APP_USER,
            "PGPASSWORD": PG_APP_PASSWORD,
            "PGDATABASE": PG_DATABASE,
        })

    build_files(dest, rng)
    report_names = build_reports(dest, f"localhost:{args.pg_port}")
    build_usage(dest, rng, report_names)
    if pg_bin is not None:
        build_postgres(dest, rng, pg_bin, args.pg_port)
    build_governance_db(dest, rng)
    bootstrap_app_state(dest)

    (dest / "sandbox_config.json").write_text(json.dumps({
        "created_at": datetime.now(timezone.utc).isoformat(),
        "pg_enabled": pg_bin is not None,
        "pg_port": args.pg_port,
        "pg_user": PG_APP_USER,
        "pg_password": PG_APP_PASSWORD,
        "pg_database": PG_DATABASE,
        "pg_bin": str(pg_bin) if pg_bin else None,
    }, indent=2), encoding="utf-8")
    (dest / "README.txt").write_text(
        "Metronome offline sandbox - generated by tools/seed_sandbox.py.\n"
        "Everything in this folder is fake, disposable test data.\n\n"
        "Start the app against it:  python tools/run_sandbox.py\n"
        "Stop the sandbox database: python tools/run_sandbox.py --stop\n"
        "Then simply delete this folder to remove every trace.\n",
        encoding="utf-8",
    )
    log(f"Sandbox ready at {dest}")
    log("Next: python tools/run_sandbox.py")


if __name__ == "__main__":
    main()

"""Synthetic Excel files: explicit names, no inferred data exclusion, no SQL."""
import csv
from pathlib import Path

import pytest
from openpyxl import Workbook

from app import flow_excel, flow_worker


def workbook_file(tmp_path, sheets):
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets:
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    source = tmp_path / "regional-orders.xlsx"
    book.save(source)
    book.close()
    return source


def normalize(source, tmp_path, config=None, **kwargs):
    return flow_worker._store_completed_download(
        source, tmp_path / "saved.xlsx", excel_worksheets=config,
        csv_preamble="none", xlsx_header_mode="first_row", strict_headers=True,
        **kwargs,
    )


def data_rows(metadata):
    with Path(metadata["file_path"]).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def test_single_sheet_default_retains_small_data_and_total_labels(tmp_path):
    source = workbook_file(tmp_path, [("Totals", [["Code", "Total"], ["A", 2], ["A", 2]])])
    result = normalize(source, tmp_path)
    assert data_rows(result) == [["Code", "Total"], ["A", "2"], ["A", "2"]]
    assert result["source_sheets"] == ["Totals"]
    assert result["sheet_row_counts"] == {"Totals": 2}
    assert Path(result["original_file_path"]).read_bytes() == source.read_bytes()


@pytest.mark.parametrize("second_rows", [[], [["Code", "Units"], ["B", 4]]])
@pytest.mark.parametrize("raw_only", [False, True])
def test_multiple_sheets_fail_by_default_even_if_empty_hidden_or_compatible(tmp_path, second_rows, raw_only):
    source = workbook_file(tmp_path, [("North", [["Code", "Units"], ["A", 3]]), ("South", second_rows)])
    from openpyxl import load_workbook
    book = load_workbook(source)
    book["South"].sheet_state = "hidden"
    book.save(source)
    book.close()
    original = source.read_bytes()
    events = []
    with pytest.raises(flow_excel.WorksheetError) as error:
        normalize(source, tmp_path, require_normalized_csv=not raw_only, recorded_output=raw_only,
                  processing_progress=lambda stage, message: events.append((stage, message)))
    assert str(error.value).startswith("This Excel has more than one sheet. Please enable the option in Flows.")
    assert "North" in str(error.value) and "South" in str(error.value)
    assert "SQL was not started" in str(error.value)
    assert events[-1][0] == "file_normalization_failed"
    assert events[-1][1] == str(error.value)
    assert source.read_bytes() == original
    assert (tmp_path / "saved.xlsx").read_bytes() == original
    assert not (tmp_path / "saved_normalized.csv").exists()


def test_append_uses_exact_supplied_order_and_retains_every_selected_row(tmp_path):
    source = workbook_file(tmp_path, [
        ("North", [["Code", "Units"], ["N", 4], ["N", 4]]),
        ("South", [["Code", "Units"], ["S", 7]]),
        ("etc.", [["Year", "Quarter", "Total"], [2026, 1, 15]]),
    ])
    events = []
    result = normalize(source, tmp_path, {"mode": "append", "names": ["South", "North"]},
                       processing_progress=lambda stage, message: events.append((stage, message)))
    assert data_rows(result) == [["Code", "Units"], ["S", "7"], ["N", "4"], ["N", "4"]]
    assert result["source_sheets"] == ["South", "North"]
    assert result["sheet_row_counts"] == {"South": 1, "North": 2}
    assert result["available_sheets"] == ["North", "South", "etc."]
    assert any("'South': 1 rows, 2 columns" in message for _, message in events)


def test_single_named_sheet_can_be_the_small_total_sheet(tmp_path):
    source = workbook_file(tmp_path, [
        ("North", [["Code", "Units"], ["N", 4]]),
        ("etc.", [["Year", "Quarter", "Total"]] + [[2026, i, i] for i in range(17)]),
    ])
    result = normalize(source, tmp_path, {"mode": "single", "names": ["etc."]})
    assert result["row_count"] == 17
    assert result["source_sheets"] == ["etc."]
    assert len(data_rows(result)) == 18


@pytest.mark.parametrize("second_header", [["Code", "Units", "Total"], ["Units", "Code"], ["Code", "Amount"]])
def test_append_mismatch_fails_with_file_sheets_and_columns_without_csv_fallback(tmp_path, second_header):
    source = workbook_file(tmp_path, [
        ("North", [["Code", "Units"], ["N", 4]]),
        ("South", [second_header, [1] * len(second_header)]),
    ])
    original = source.read_bytes()
    with pytest.raises(flow_excel.WorksheetError) as error:
        normalize(source, tmp_path, {"mode": "append", "names": ["North", "South"]}, allow_raw_xlsx_fallback=True)
    assert error.value.details["code"] == "excel_columns_mismatch"
    for expected in ["saved.xlsx", "North", "South", "different columns", "Expected:", "Found:", "SQL was not started"]:
        assert expected in str(error.value)
    assert not (tmp_path / "saved_normalized.csv").exists()
    assert (tmp_path / ".saved_normalized.csv.partial").is_file()
    assert (tmp_path / "saved.xlsx").read_bytes() == original


def test_missing_or_renamed_sheet_fails_before_publishing(tmp_path):
    source = workbook_file(tmp_path, [("North 2027", [["Code", "Units"], ["N", 4]])])
    with pytest.raises(flow_excel.WorksheetError) as error:
        normalize(source, tmp_path, {"mode": "single", "names": ["North"]})
    assert error.value.details["code"] == "excel_sheet_missing"
    assert "North 2027" in str(error.value)
    assert not (tmp_path / "saved_normalized.csv").exists()


def test_explicitly_selected_empty_sheet_is_not_silently_skipped(tmp_path):
    source = workbook_file(tmp_path, [("North", [["Code", "Units"], ["N", 4]]), ("South", [])])
    with pytest.raises(flow_excel.WorksheetError, match="South.*did not contain a usable table"):
        normalize(source, tmp_path, {"mode": "append", "names": ["North", "South"]})
    assert not (tmp_path / "saved_normalized.csv").exists()


@pytest.mark.parametrize("config", [
    {"mode": "auto", "names": ["A"]}, {"mode": "single", "names": []},
    {"mode": "single", "names": ["A", "B"]}, {"mode": "append", "names": ["A"]},
    {"mode": "append", "names": ["A", "A"]}, {"mode": "single", "names": [" "]},
    {"mode": "single", "names": [1]}, {"mode": "single", "names": ["A"], "skip_small": True},
    {"mode": [], "names": ["A"]},
])
def test_invalid_choices_fail_instead_of_guessing(config):
    with pytest.raises(ValueError):
        flow_excel.normalize_config(config)


def test_selection_is_case_sensitive_and_preserves_exact_name():
    with pytest.raises(flow_excel.WorksheetError):
        flow_excel.select_names(["North"], {"mode": "single", "names": ["north"]}, workbook="orders.xlsx")
    assert flow_excel.select_names([" North "], {"mode": "single", "names": [" North "]}, workbook="orders.xlsx") == [" North "]


@pytest.mark.parametrize("payload", [b"Code,Units\nN,4\n", b"<html><table><tr><th>Code</th><th>Units</th></tr><tr><td>N</td><td>4</td></tr></table></html>"])
def test_named_worksheet_option_cannot_silently_load_a_non_workbook(tmp_path, payload):
    source = tmp_path / "changed-download.xlsx"
    source.write_bytes(payload)
    with pytest.raises(flow_excel.WorksheetError) as error:
        normalize(source, tmp_path, {"mode": "single", "names": ["North"]})
    assert error.value.details["code"] == "excel_workbook_required"
    assert source.read_bytes() == payload
    assert not (tmp_path / "saved_normalized.csv").exists()


def test_wrapped_portal_failure_keeps_structured_worksheet_evidence():
    try:
        try:
            flow_excel.select_names(["North", "South"], None, workbook="orders.xlsx")
        except flow_excel.WorksheetError as exc:
            raise flow_worker._CompletedDownloadProcessingError(str(exc)) from exc
    except flow_worker._CompletedDownloadProcessingError as error:
        detail = flow_excel.failure_details(error)
    assert detail["stage"] == "file_normalization_failed"
    assert detail["excel"]["available_sheets"] == ["North", "South"]
    assert detail["excel"]["workbook"] == "orders.xlsx"
    assert detail["excel"]["sql_started"] is False
    assert flow_excel.failure_details(RuntimeError("Unrelated")) == {}


@pytest.mark.parametrize("choice", [None, {"mode": "append", "names": ["South", "North"]}])
def test_generic_portal_download_uses_the_same_worksheet_rule(tmp_path, monkeypatch, choice):
    from types import SimpleNamespace
    from unittest.mock import MagicMock
    source = workbook_file(tmp_path, [("North", [["Code", "Units"], ["N", 4]]),
                                     ("South", [["Code", "Units"], ["S", 7]])])
    download = SimpleNamespace(save_as=lambda path: Path(path).write_bytes(source.read_bytes()),
                               suggested_filename=source.name)
    page = MagicMock()
    page.expect_download.return_value.__enter__.return_value = SimpleNamespace(value=download)
    monkeypatch.setattr(flow_worker, "_click_named", lambda *args, **kwargs: None)
    (tmp_path / "output").mkdir()
    job = {
        "flow": {"id": 1, "name": "Orders"}, "site": {"adapter": "web_export"},
        "report": {"id": 1, "name": "Orders", "url": "https://portal.invalid/orders", "download_text": "Export", "filters": []},
        "selections": {}, "downloads": {"target_folder": str(tmp_path / "output"),
            "file_format": "xlsx", "filename_template": "orders.xlsx", "periods": [None], "excel_worksheets": choice},
    }
    def run():
        return flow_worker.execute_job(page, job, lambda *args, **kwargs: None,
            tmp_path / "profile", run_id=92, register_folder=lambda path: {"ops": []})
    if choice is None:
        with pytest.raises(flow_worker._CompletedDownloadProcessingError, match="more than one sheet"):
            run()
    else:
        artifacts, _ = run()
        assert data_rows(artifacts[0]) == [["Code", "Units"], ["S", "7"], ["N", "4"]]


def test_empty_data_failure_keeps_workbook_and_sheet_in_log(tmp_path):
    source = workbook_file(tmp_path, [("North", [["Code", "Units"]])])
    events = []
    with pytest.raises(flow_excel.WorksheetError) as error:
        normalize(source, tmp_path, processing_progress=lambda stage, message: events.append((stage, message)))
    assert "North" in str(error.value)
    assert "SQL was not started" in str(error.value)
    assert events[-1][0] == "file_normalization_failed"
